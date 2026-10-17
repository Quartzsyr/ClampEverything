"""Author the real product files used by the interactive desktop. No personal inputs."""
from pathlib import Path
import json,shutil,zipfile
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];files=root/'assets/files';source=root/'media-source/onboarding'
source.mkdir(exist_ok=True)
note='''随手夹取的小提示

按住 Option，鼠标变成夹子。
瞄准内容，按下并拖到需要的位置。
靠边收起，悬停展开。

点纸张复制，拖出原文件继续用。
试试向助手发送“怎么用”或“发我指南”。
'''
(files/'夹取操作提示.txt').write_text(note)
code='''// QuickNote.swift - 在 macOS SwiftUI 项目中运行的最小便笺示例。
// 将此文件设为项目入口；网页编辑器仅编辑文件，不执行 Swift。
import SwiftUI

@main
struct QuickNoteApp: App {
    var body: some Scene {
        WindowGroup {
            QuickNoteView()
        }
        .defaultSize(width: 320, height: 240)
    }
}

struct QuickNoteView: View {
    @AppStorage("quick-note-text") private var text = "把灵感，夹在手边。"

    var body: some View {
        VStack(alignment: .leading, spacing: 12) {
            Label("我的便笺", systemImage: "note.text")
                .font(.headline)
            TextEditor(text: $text)
                .font(.body)
                .scrollContentBackground(.hidden)
            Text("文字保存在这台 Mac 上。")
                .font(.caption)
                .foregroundStyle(.secondary)
        }
        .padding(20)
        .frame(minWidth: 280, minHeight: 200)
        .background(Color(red: 0.98, green: 0.97, blue: 0.91))
    }
}
'''
(files/'QuickNote.swift').write_text(code)
data={'说明':'这是一份可阅读的演示清单，不是应用导入格式，也不包含个人数据。','场景':'阅读与写作','资料':[{'名称':'快速上手指南','类型':'PDF','用途':'放在左侧边缘参考'},{'名称':'我的桌面收纳计划','类型':'Word','用途':'记录今天想整理的内容'},{'名称':'山间湖泊','类型':'照片','用途':'选择喜欢的画面'}],'练习':['夹取一段文字','展开 PDF 并翻到下一页','修改 Word 或代码并保存','拖到废纸篓后撤销'],'官网':'https://clamp.musefilm.top/'}
(files/'桌面整理示例.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
voice='''欢迎使用屏幕夹子。按住 Option 键，鼠标会变成夹子。瞄准文字、照片或文件，按下鼠标，拖到需要的位置，再松开。靠近屏幕边缘可以收起，鼠标停留后展开。打开快速上手指南，试试翻页。也可以编辑收纳计划，或把代码改成自己的便笺。喜欢的话，在官网下载 Mac 版。'''
(source/'操作讲解文字稿.txt').write_text('系统合成语音，不是真人录音。\n\n'+voice+'\n')
# PDF: three small, readable pages using the actual recorded desktop screenshots.
pdfmetrics.registerFont(TTFont('DemoChinese','/System/Library/Fonts/Supplemental/Arial Unicode.ttf'))
w,h=420,595;c=canvas.Canvas(str(files/'屏幕夹子快速上手.pdf'),pagesize=(w,h));c.setTitle('屏幕夹子快速上手');c.setAuthor('Quartz')
def text(x,y,s,size=12,color='#344335'):
 c.setFillColor(color);c.setFont('DemoChinese',size);c.drawString(x,y,s)
def base(n,title,sub):
 c.setFillColorRGB(.98,.98,.95);c.rect(0,0,w,h,fill=1,stroke=0)
 text(30,555,'CLAMP EVERYTHING / QUICK START',9,'#78906b')
 text(30,510,title,27);text(30,479,sub,11,'#72806b')
 text(30,25,'clamp.musefilm.top',9,'#72806b');text(370,25,f'0{n}',10,'#72806b')
def image(name,y,height):
 p=root/'assets/usage'/name;im=Image.open(p);iw,ih=im.size;scale=min(360/iw,height/ih);dw,dh=iw*scale,ih*scale
 c.drawImage(ImageReader(im),30+(360-dw)/2,y+height-dh,width=dw,height=dh,mask='auto')
base(1,'把内容夹在手边','先试一次文字夹取，再试试桌面的真实文件。');image('text-pickup.webp',224,225)
for y,s in [(194,'01  按住 Option，鼠标变成夹子。'),(167,'02  瞄准文字、照片或文件，按下鼠标。'),(140,'03  拖到需要的位置，松开即可放置。'),(94,'网页中，普通鼠标单击选择、双击打开文件。'),(73,'Mac 应用请以设置中的夹取键与系统权限为准。')]:text(30,y,s,11)
c.showPage()
base(2,'展开阅读 随手取用','文稿保留原文件，阅读和复制可以接着进行。');image('documents.webp',224,225)
for y,s in [(194,'PDF 与 Word：用滚轮或页码按钮翻页。'),(167,'点击夹住的内容：复制到网页桌面的剪贴板。'),(140,'拖出夹项：继续使用对应的真实文件。'),(98,'编辑练习：打开 Word 或 Swift 文件，修改并保存。'),(76,'网页中的文件复制与系统剪贴板受浏览器能力限制。')]:text(30,y,s,11)
c.showPage()
base(3,'收起资料 留出空间','让资料留在身边，也让桌面保持轻松。');image('edge-unfold.webp',224,225)
for y,s in [(194,'靠近屏幕左右边缘，资料会折叠收起。'),(167,'鼠标停留在夹子上，展开查看内容。'),(140,'拖住夹子可以移动，拖出内容用于复制。'),(98,'网页练习：把图标拖到废纸篓，再撤销恢复。'),(76,'更多说明与 Mac 版下载：clamp.musefilm.top')]:text(30,y,s,11)
c.save()
# Word: deliberately editable three-page plan, without decorative boxes.
doc=Document();
for element in list(doc.styles.element.iter()):
 if element.tag==qn('w:pBdr'):
  element.getparent().remove(element)
sec=doc.sections[0];sec.page_width=Cm(14.8);sec.page_height=Cm(21);sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.6);sec.right_margin=Cm(1.6)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='PingFang SC';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(13 if name=='Normal' else 25 if name=='Title' else 20)
 st.element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC')
 st.paragraph_format.space_after=Pt(10);st.paragraph_format.line_spacing=1.4

def para(s,style=None):
 p=doc.add_paragraph(style=style);r=p.add_run(s);r.font.name='PingFang SC';r.font.size=Pt(25 if style=='Title' else 20 if style=='Heading 1' else 13);r.font.color.rgb=RGBColor(0,0,0);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC');return p
para('我的桌面收纳计划','Title');para('这份计划用于练习编辑和整理资料。把示例改成自己的目标，保存后再夹在手边。')
para('今天要留住什么','Heading 1')
for s in ['我的工作场景：阅读与写作','想留下的一句话：把灵感夹在手边','需要随时参考的文稿：快速上手指南','喜欢的照片：山间湖泊','完成后：保存这份计划，拖出一份副本。']:para(s)
doc.add_page_break();para('为资料安排位置','Heading 1');para('让经常查看的内容靠边，让正在编辑的内容留在中央。')
for s in ['左侧边缘：参考 PDF 与摘录文字','右侧边缘：照片和临时资料','中央区域：正在修改的 Word 文稿','Dock：常用应用和文件快捷入口','练习：把一张照片夹到左侧，再展开查看。']:para(s)
doc.add_page_break();para('完成一次整理练习','Heading 1');para('这些项目都可以在网页桌面中尝试。改完后，点击保存或使用 Command S。')
for s in ['1  夹取一句话，并展开阅读。','2  打开 PDF，翻到第二页。','3  修改这份计划，保存成实际文件。','4  编辑 Swift 示例中的便笺文字。','5  拖动图标到 Dock，再尝试移除。','6  将示例文件移到废纸篓，再撤销。','我的下一步：________________']:para(s)
doc.core_properties.title='我的桌面收纳计划';doc.core_properties.author='Quartz';doc.core_properties.last_modified_by='Quartz';doc.save(files/'我的桌面收纳计划.docx')
print('Created product PDF, editable DOCX, TXT, runnable Swift example, JSON and narration script.')
# Rebuild the two downloads from the same real files. The material pack adds a transcript.
folder=files/'屏幕夹子入门资料';folder.mkdir(exist_ok=True)
for name in ['夹取操作提示.txt','屏幕夹子快速上手.pdf','我的桌面收纳计划.docx','QuickNote.swift','桌面整理示例.json','山间湖泊.png','云下牧场.jpg','夹取操作讲解.wav']:
 if (files/name).exists():shutil.copyfile(files/name,folder/name)
for name in ['屏幕夹子入门资料.zip','屏幕夹子演示素材包.zip']:
 with zipfile.ZipFile(files/name,'w',zipfile.ZIP_DEFLATED) as archive:
  for path in sorted(folder.iterdir()):
   if path.is_file():archive.write(path,folder.name+'/'+path.name)
  if name=='屏幕夹子演示素材包.zip':archive.write(source/'操作讲解文字稿.txt',folder.name+'/操作讲解文字稿.txt')
import fitz,hashlib
pdf=files/'屏幕夹子快速上手.pdf';pages=[]
for page in fitz.open(pdf):
 lines=[]
 for block in page.get_text('dict')['blocks']:
  for line in block.get('lines',[]):
   content=''.join(span['text'] for span in line['spans']);x0,y0,x1,y1=line['bbox']
   if content.strip():lines.append(dict(text=content,x=x0/page.rect.width,y=y0/page.rect.height,width=(x1-x0)/page.rect.width,height=(y1-y0)/page.rect.height))
 pages.append(lines)
pdf.with_suffix('.text.json').write_text(json.dumps({'method':'Exact text and bounding boxes extracted from the matching authored PDF','sourceSHA256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'pages':pages},ensure_ascii=False,indent=2)+'\n')
